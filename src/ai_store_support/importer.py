from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from .normalization import normalize

HEADER_ALIASES = {
    "material": ("材质", "材质名称", "material", "material_name"),
    "brand": ("品牌栏目", "品牌", "brand"),
    "model_group": ("型号组合", "型号", "机型", "model", "model_group"),
    "stock_status": ("库存状态", "状态", "stock", "stock_status"),
}


def _match_headers(headers: list[str]) -> dict[str, int]:
    normalized = [normalize(item) for item in headers]
    result: dict[str, int] = {}
    for key, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            target = normalize(alias)
            if target in normalized:
                result[key] = normalized.index(target)
                break
    if "model_group" not in result:
        raise ValueError("未找到‘型号组合/型号’列")
    return result


def iter_xlsx_rows(path: Path, sheet_name: str = "型号明细") -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    index_map: dict[str, int] | None = None
    for row in rows:
        values = ["" if value is None else str(value).strip() for value in row]
        if index_map is None:
            try:
                index_map = _match_headers(values)
            except ValueError:
                continue
            continue
        get = lambda key: values[index_map[key]] if key in index_map and index_map[key] < len(values) else ""
        yield {
            "material": get("material"),
            "brand": get("brand"),
            "model_group": get("model_group"),
            "stock_status": get("stock_status"),
        }


def iter_csv_rows(path: Path) -> Iterable[dict[str, Any]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise ValueError("无法识别CSV编码")
    reader = csv.reader(text.splitlines())
    index_map: dict[str, int] | None = None
    for row in reader:
        values = [str(value).strip() for value in row]
        if index_map is None:
            try:
                index_map = _match_headers(values)
            except ValueError:
                continue
            continue
        get = lambda key: values[index_map[key]] if key in index_map and index_map[key] < len(values) else ""
        yield {
            "material": get("material"),
            "brand": get("brand"),
            "model_group": get("model_group"),
            "stock_status": get("stock_status"),
        }


def iter_source_rows(path: str | Path) -> Iterable[dict[str, Any]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return iter_xlsx_rows(source)
    if suffix == ".csv":
        return iter_csv_rows(source)
    raise ValueError(f"暂不支持的文件格式: {suffix}")
