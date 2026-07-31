from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Callable, Iterable

from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from .config_models import KnowledgeEntry
from .normalization import normalize
from .schemas import KnowledgeMatch

KNOWLEDGE_HEADER_ALIASES = {
    "title": ("标题", "问题", "知识标题", "title", "question"),
    "content": ("内容", "答案", "回复", "知识内容", "content", "answer"),
    "tags": ("标签", "关键词", "tags", "keywords"),
    "priority": ("优先级", "priority"),
    "enabled": ("启用", "是否启用", "enabled"),
    "direct_reply": ("直接回复", "是否直接回复", "direct_reply"),
}


def _as_bool(value: Any, default: bool = True) -> bool:
    text = str(value or "").strip().lower()
    if not text:
        return default
    return text not in {"0", "false", "否", "不启用", "关闭", "no"}


def _split_tags(value: Any) -> list[str]:
    return [item.strip() for item in re.split(r"[,，、|｜;；]+", str(value or "")) if item.strip()]


def _match_headers(headers: list[str]) -> dict[str, int]:
    normalized_headers = [normalize(item) for item in headers]
    result: dict[str, int] = {}
    for field, aliases in KNOWLEDGE_HEADER_ALIASES.items():
        for alias in aliases:
            key = normalize(alias)
            if key in normalized_headers:
                result[field] = normalized_headers.index(key)
                break
    if "title" not in result or "content" not in result:
        raise ValueError("知识库文件必须包含标题/问题列和内容/答案列")
    return result


def _iter_csv(path: Path) -> Iterable[dict[str, Any]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别知识库CSV编码")
    rows = csv.reader(text.splitlines())
    index_map: dict[str, int] | None = None
    for row in rows:
        values = [str(value or "").strip() for value in row]
        if index_map is None:
            try:
                index_map = _match_headers(values)
            except ValueError:
                continue
            continue
        yield _values_to_entry(values, index_map)


def _iter_xlsx(path: Path) -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    index_map: dict[str, int] | None = None
    for row in worksheet.iter_rows(values_only=True):
        values = [str(value or "").strip() for value in row]
        if index_map is None:
            try:
                index_map = _match_headers(values)
            except ValueError:
                continue
            continue
        yield _values_to_entry(values, index_map)


def _values_to_entry(values: list[str], index_map: dict[str, int]) -> dict[str, Any]:
    def get(field: str) -> str:
        index = index_map.get(field)
        return values[index] if index is not None and index < len(values) else ""

    priority_text = get("priority")
    try:
        priority = int(float(priority_text)) if priority_text else 0
    except ValueError:
        priority = 0
    return {
        "title": get("title"),
        "content": get("content"),
        "tags": _split_tags(get("tags")),
        "priority": priority,
        "enabled": _as_bool(get("enabled"), True),
        "direct_reply": _as_bool(get("direct_reply"), True),
    }


def iter_knowledge_rows(path: str | Path) -> Iterable[dict[str, Any]]:
    source = Path(path)
    if source.suffix.lower() == ".csv":
        return _iter_csv(source)
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        return _iter_xlsx(source)
    raise ValueError(f"暂不支持的知识库格式: {source.suffix}")


class KnowledgeService:
    def __init__(self, session_factory: Callable[[], Session]):
        self.session_factory = session_factory

    def upsert(
        self,
        *,
        shop_key: str,
        title: str,
        content: str,
        tags: list[str] | tuple[str, ...] = (),
        priority: int = 0,
        enabled: bool = True,
        direct_reply: bool = True,
    ) -> int:
        title = str(title or "").strip()
        content = str(content or "").strip()
        if not title or not content:
            raise ValueError("知识标题和内容不能为空")
        with self.session_factory() as session:
            row = session.scalar(
                select(KnowledgeEntry).where(
                    KnowledgeEntry.shop_key == shop_key,
                    KnowledgeEntry.title == title,
                )
            )
            if row is None:
                row = KnowledgeEntry(shop_key=shop_key, title=title, content=content)
                session.add(row)
            row.content = content
            row.tags_json = json.dumps(list(tags), ensure_ascii=False)
            row.priority = int(priority)
            row.enabled = bool(enabled)
            row.direct_reply = bool(direct_reply)
            session.commit()
            return row.id

    def import_file(self, shop_key: str, path: str | Path) -> dict[str, int]:
        stats = {"rows": 0, "created_or_updated": 0, "skipped": 0}
        for item in iter_knowledge_rows(path):
            stats["rows"] += 1
            if not item["title"] or not item["content"]:
                stats["skipped"] += 1
                continue
            self.upsert(shop_key=shop_key, **item)
            stats["created_or_updated"] += 1
        return stats

    def search(self, shop_key: str, query: str, limit: int = 5) -> list[KnowledgeMatch]:
        query = str(query or "").strip()
        if not query:
            return []
        with self.session_factory() as session:
            rows = list(
                session.scalars(
                    select(KnowledgeEntry)
                    .where(
                        KnowledgeEntry.shop_key == shop_key,
                        KnowledgeEntry.enabled.is_(True),
                    )
                    .order_by(desc(KnowledgeEntry.priority), KnowledgeEntry.id)
                )
            )
        scored: list[KnowledgeMatch] = []
        query_norm = normalize(query)
        query_terms = self._query_terms(query)
        for row in rows:
            tags = tuple(json.loads(row.tags_json or "[]"))
            score = self._score(query_norm, query_terms, row.title, row.content, tags)
            if score <= 0:
                continue
            scored.append(
                KnowledgeMatch(
                    entry_id=row.id,
                    title=row.title,
                    content=row.content,
                    tags=tags,
                    score=score + row.priority * 0.01,
                    direct_reply=row.direct_reply,
                )
            )
        return sorted(scored, key=lambda item: (-item.score, item.entry_id))[:limit]

    @staticmethod
    def _query_terms(query: str) -> set[str]:
        normalized = normalize(query)
        terms = {normalize(item) for item in re.findall(r"[A-Za-z]+\d*|\d+[A-Za-z]*|[\u4e00-\u9fff]{2,}", query)}
        if len(normalized) >= 2:
            terms.update(normalized[index : index + 2] for index in range(len(normalized) - 1))
        terms.discard("")
        return terms

    @staticmethod
    def _score(
        query_norm: str,
        query_terms: set[str],
        title: str,
        content: str,
        tags: tuple[str, ...],
    ) -> float:
        title_norm = normalize(title)
        content_norm = normalize(content)
        tag_norms = [normalize(tag) for tag in tags if normalize(tag)]
        score = 0.0
        if query_norm and query_norm == title_norm:
            score += 20.0
        elif query_norm and (query_norm in title_norm or title_norm in query_norm):
            score += 10.0
        for tag in tag_norms:
            if tag == query_norm:
                score += 8.0
            elif tag and tag in query_norm:
                score += 5.0
        title_hits = sum(1 for term in query_terms if len(term) >= 2 and term in title_norm)
        tag_hits = sum(1 for term in query_terms if any(term in tag for tag in tag_norms))
        content_hits = sum(1 for term in query_terms if len(term) >= 2 and term in content_norm)
        score += min(title_hits, 6) * 2.0
        score += min(tag_hits, 6) * 1.5
        score += min(content_hits, 6) * 0.35
        return score
